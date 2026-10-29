"""The whole app in a headless browser, with the stand-in AI. Run `pytest -m "not ui"` to skip these."""

import re
import threading
import time

import openpyxl
import pytest

from ai.types import AIQuotaError
from tests.conftest import SAMPLES
from tests.fake_llm import FakeProvider
from utils import components as ui
from utils import constants as C
from utils import session

playwright_sync = pytest.importorskip("playwright.sync_api")
expect = playwright_sync.expect
pytestmark = pytest.mark.ui

GREEN, TRADITIONAL = SAMPLES / "sample_budget_green_focus.pdf", SAMPLES / "sample_budget_traditional.pdf"
PROJECT_DOCX = SAMPLES / "sample_project_budget.docx"
SLOW = 60_000
TECHNICAL_WORDS = re.compile(r"\b(gemini|llm|api|json|prompt|databricks|token|tokens|python|dash|plotly|pandas|gpt|openai|anthropic)\b", re.I)


@pytest.fixture(scope="module")
def base():
    from werkzeug.serving import make_server
    import app as app_module
    server = make_server("127.0.0.1", 0, app_module.server, threaded=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield "http://127.0.0.1:%d" % server.server_port
    server.shutdown()


@pytest.fixture(scope="module")
def browser():
    with playwright_sync.sync_playwright() as p:
        b = p.chromium.launch()
        yield b
        b.close()


@pytest.fixture
def page(browser):
    ctx = browser.new_context(viewport={"width": 1360, "height": 900}, accept_downloads=True)
    pg = ctx.new_page()
    pg.problems = []
    pg.on("pageerror", lambda e: pg.problems.append("pageerror: %s" % e))
    pg.on("console", lambda m: pg.problems.append("console: %s" % m.text) if m.type == "error" and "favicon" not in m.text else None)
    yield pg
    ctx.close()


def go(page, base, name):
    page.goto("%s/?page=%s" % (base, name))
    page.wait_for_selector(".page")


def drop(page, selector, path):
    page.set_input_files("%s input[type=file]" % selector, str(path))


def charts_drawn(page, minimum):
    page.wait_for_function("n => document.querySelectorAll('.results .js-plotly-plot').length >= n", arg=minimum, timeout=SLOW)


def run_evaluate(page, base, path=GREEN):
    go(page, base, "evaluate")
    drop(page, "#eval-upload", path)
    page.wait_for_selector(".results", timeout=SLOW)
    charts_drawn(page, 1)


def run_compare(page, base):
    go(page, base, "compare")
    drop(page, "#compare-upload-a", GREEN)
    page.wait_for_selector("#compare-file-a .file-chip")
    drop(page, "#compare-upload-b", TRADITIONAL)
    page.wait_for_selector("#compare-file-b .file-chip")
    page.click("#compare-run")
    page.wait_for_selector(".results", timeout=SLOW)
    charts_drawn(page, 1)


def visible_text(page):
    return page.inner_text("body")


# ── pages and navigation ────────────────────────────────────────────────────
def test_home_offers_the_two_tools(page, base):
    go(page, base, "home")
    assert page.title() == "Home - " + C.APP_NAME
    assert [a.inner_text().strip() for a in page.query_selector_all(".hero-actions a")] == ["Evaluate Budget", "Compare Budget", "How it works"]
    assert page.is_visible(".brand-mark") and "World Bank" in page.get_attribute(".footer-wb", "alt")
    assert "smarter budgets" in page.inner_text(".site-footer")
    assert page.problems == []


def test_the_menu_has_only_the_two_tools_besides_home_and_methodology(page, base):
    go(page, base, "home")
    labels = [a.inner_text().strip() for a in page.query_selector_all(".header-nav a")]
    assert labels == ["Home", "Evaluate Budget", "Compare Budget", "Methodology"]


def test_each_menu_item_opens_its_page(page, base):
    go(page, base, "home")
    page.click(".header-nav a:has-text('Evaluate Budget')")
    page.wait_for_selector("#eval-upload")
    assert "page=evaluate" in page.url and page.title().startswith("Evaluate Budget")
    page.click(".header-nav a:has-text('Compare Budget')")
    page.wait_for_selector("#compare-upload-a")
    assert "page=compare" in page.url
    page.click(".header-brand")
    page.wait_for_selector(".hero")
    assert "page=home" in page.url and page.problems == []


def test_methodology_opens_in_a_new_tab_and_is_plain_language(page, base):
    go(page, base, "home")
    with page.context.expect_page() as opened:
        page.click(".header-nav a:has-text('Methodology')")
    guide = opened.value
    guide.wait_for_selector(".doc-layout")
    assert "page=methodology" in guide.url and page.url.endswith("page=home")
    text = guide.inner_text("body")
    for heading in ("Rio Markers", "EU Taxonomy", "The overall score", "Spending we do not count"):
        assert heading in text
    assert not TECHNICAL_WORDS.search(text), TECHNICAL_WORDS.search(text).group(0)
    assert len(guide.query_selector_all(".doc-toc a")) == 10


def test_how_it_works_button_also_opens_a_new_tab(page, base):
    go(page, base, "home")
    with page.context.expect_page() as opened:
        page.click(".hero-actions a:has-text('How it works')")
    opened.value.wait_for_selector(".doc-layout")


def test_unknown_pages_say_so_kindly(page, base):
    go(page, base, "nonsense")
    assert "Page not found" in page.inner_text("main")
    page.click("main a:has-text('Back to the home page')")
    page.wait_for_selector(".hero")


# ── evaluate ────────────────────────────────────────────────────────────────
def test_evaluate_shows_both_taxonomies(page, base, fake_ai):
    run_evaluate(page, base)
    text = page.inner_text(".results")
    for expected in ("Rio Markers", "EU Taxonomy", "Principal", "Significant", "Green", "Brown", "Mitigation", "Adaptation",
                     "How each tag was decided"):
        assert expected in text
    charts_drawn(page, 5)
    assert len(page.query_selector_all(".results .kpi")) >= 6
    assert page.problems == []


def test_evaluate_table_lists_every_line_and_switches_taxonomy(page, base, fake_ai):
    run_evaluate(page, base)
    page.wait_for_selector("#eval-table-rio")
    header = page.inner_text("#eval-table-rio")
    assert "Rio marker" in header and "Decided by" in header and "Why" in header
    page.click(".segmented .seg-label:has-text('EU Taxonomy')")
    page.wait_for_selector("#eval-table-eu")
    assert "Climate tag" in page.inner_text("#eval-table-eu")


def test_evaluate_download_is_an_excel_file(page, base, fake_ai, tmp_path):
    run_evaluate(page, base)
    with page.expect_download() as dl:
        page.click("#eval-download-btn")
    target = tmp_path / "out.xlsx"
    dl.value.save_as(str(target))
    assert dl.value.suggested_filename.endswith(".xlsx")
    assert openpyxl.load_workbook(target).sheetnames == ["About", "Summary", "Rio Markers", "EU Taxonomy"]


def test_evaluate_also_reads_word_files(page, base, fake_ai):
    run_evaluate(page, base, PROJECT_DOCX)
    assert "Early warning system" in page.inner_text("#eval-table-rio") or page.query_selector("#eval-table-rio")


def test_the_loading_bar_names_each_step_and_never_shows_a_percentage(page, base, use_provider):
    use_provider(FakeProvider(delay=1.3))
    go(page, base, "evaluate")
    drop(page, "#eval-upload", GREEN)
    seen = set()
    end = time.time() + 40
    while time.time() < end and not page.query_selector(".results"):
        loader = page.query_selector(".loader")
        if loader:
            stage = loader.query_selector(".loader-stage")
            if stage:
                seen.add(stage.inner_text().strip().lower())
            assert "%" not in loader.inner_text()
        time.sleep(0.1)
    assert page.query_selector(".results")
    assert any("tagging" in s for s in seen) and len(seen) >= 3, seen
    assert all(s.endswith("…") for s in seen)


def test_the_loading_bar_keeps_moving_until_the_work_is_done(page, base, use_provider):
    use_provider(FakeProvider(delay=0.8))
    go(page, base, "evaluate")
    drop(page, "#eval-upload", GREEN)
    page.wait_for_selector(".loader-bar")
    style = page.eval_on_selector(".loader-bar", "e => { const s = getComputedStyle(e); return [s.animationName, s.animationIterationCount]; }")
    assert style[0] != "none" and style[1] == "infinite"
    x1 = page.eval_on_selector(".loader-bar", "e => e.getBoundingClientRect().left")
    time.sleep(0.4)
    x2 = page.eval_on_selector(".loader-bar", "e => e.getBoundingClientRect().left")
    assert x1 != x2


def test_stopping_an_analysis_says_so_and_keeps_nothing(page, base, use_provider):
    use_provider(FakeProvider(delay=1.0))
    go(page, base, "evaluate")
    drop(page, "#eval-upload", GREEN)
    page.wait_for_selector(".loader")
    page.click("#eval-cancel")
    page.wait_for_selector(".notice-info:has-text('stopped')", timeout=10_000)
    assert not page.is_visible("#eval-cancel") and page.query_selector(".results") is None


def test_a_wrong_file_type_is_explained(page, base, fake_ai, tmp_path):
    note = tmp_path / "notes.txt"
    note.write_text("hello")
    go(page, base, "evaluate")
    drop(page, "#eval-upload", note)
    page.wait_for_selector(".notice-error", timeout=10_000)
    assert "PDF" in page.inner_text(".notice-error") and page.query_selector(".loader") is None


def test_a_file_that_is_too_large_is_explained(page, base, fake_ai, tmp_path, monkeypatch):
    monkeypatch.setenv(C.ENV_MAX_UPLOAD_MB, "1")
    big = tmp_path / "big.pdf"
    big.write_bytes(b"%PDF-1.4 " + b"0" * (2 * 1024 * 1024))
    go(page, base, "evaluate")
    drop(page, "#eval-upload", big)
    page.wait_for_selector(".notice-error", timeout=20_000)
    assert "1 MB" in page.inner_text(".notice-error") and fake_ai.calls == []


def test_a_broken_pdf_is_explained(page, base, fake_ai, tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 nothing here")
    go(page, base, "evaluate")
    drop(page, "#eval-upload", bad)
    page.wait_for_selector(".notice-error", timeout=20_000)
    assert "Traceback" not in page.inner_text("main")


def test_ai_trouble_shows_a_plain_message(page, base, use_provider):
    use_provider(FakeProvider(fail_first=99, error=AIQuotaError("RESOURCE_EXHAUSTED 429")))
    go(page, base, "evaluate")
    drop(page, "#eval-upload", GREEN)
    page.wait_for_selector(".notice-error", timeout=30_000)
    text = page.inner_text(".notice-error")
    assert "daily limit" in text and "RESOURCE_EXHAUSTED" not in text and "429" not in text


def test_a_missing_key_is_explained_before_any_work(page, base, use_provider):
    class Unset(FakeProvider):
        def is_configured(self):
            return False

    provider = use_provider(Unset())
    go(page, base, "evaluate")
    drop(page, "#eval-upload", GREEN)
    page.wait_for_selector(".notice-error:has-text('not set up')", timeout=10_000)
    assert provider.calls == [] and "access key" in page.inner_text(".notice-error")


def test_results_survive_going_to_another_page_and_back(page, base, fake_ai):
    run_evaluate(page, base)
    page.click(".header-nav a:has-text('Home')")
    page.wait_for_selector(".hero")
    page.click(".hero-actions a:has-text('Evaluate Budget')")
    page.wait_for_selector(".results", timeout=15_000)


# ── compare ─────────────────────────────────────────────────────────────────
def test_compare_puts_two_budgets_side_by_side_and_ends_with_the_score(page, base, fake_ai):
    run_compare(page, base)
    rows = page.query_selector_all(".compare-table tbody tr")
    last = [c.inner_text().strip() for c in rows[-1].query_selector_all("th, td")]
    assert last[0].startswith("Overall climate score")
    scores = [int(re.search(r"\d+", c).group(0)) for c in last[1:]]
    assert len(scores) == 2 and scores[0] > scores[1] and all(0 <= s <= 100 for s in scores)
    heads = page.inner_text(".compare-table thead")
    assert "sample_budget_green_focus.pdf" in heads and "sample_budget_traditional.pdf" in heads
    assert [m.inner_text() for m in page.query_selector_all(".compare-table thead .slot-mark")] == ["A", "B"]
    charts_drawn(page, 2)
    assert page.problems == []


def test_compare_download_has_the_comparison(page, base, fake_ai, tmp_path):
    run_compare(page, base)
    with page.expect_download() as dl:
        page.click("#compare-download-btn")
    target = tmp_path / "cmp.xlsx"
    dl.value.save_as(str(target))
    assert openpyxl.load_workbook(target).sheetnames == ["Comparison", "Scores"]


def test_compare_asks_for_both_files(page, base, fake_ai):
    go(page, base, "compare")
    page.click("#compare-run")
    expect(page.locator("#compare-message")).to_contain_text("both budgets", timeout=10_000)
    drop(page, "#compare-upload-a", GREEN)
    page.wait_for_selector("#compare-file-a .file-chip")
    page.click("#compare-run")
    expect(page.locator("#compare-message")).to_contain_text("budget B", timeout=10_000)
    assert page.query_selector(".results") is None


def test_each_side_of_compare_is_independent_of_evaluate(page, base, fake_ai):
    run_evaluate(page, base)
    go(page, base, "compare")
    assert page.query_selector("#compare-file-a .file-chip") is None
    assert page.query_selector(".results") is None


def test_compare_shows_progress_for_each_budget(page, base, use_provider):
    use_provider(FakeProvider(delay=0.5))
    go(page, base, "compare")
    drop(page, "#compare-upload-a", GREEN)
    page.wait_for_selector("#compare-file-a .file-chip")
    drop(page, "#compare-upload-b", TRADITIONAL)
    page.wait_for_selector("#compare-file-b .file-chip")
    page.click("#compare-run")
    seen = set()
    end = time.time() + 60
    while time.time() < end and not page.query_selector(".results"):
        el = page.query_selector(".loader-detail")
        if el:
            seen.add(el.inner_text().strip()[:8])
        time.sleep(0.1)
    assert any(s.startswith("Budget A") for s in seen) and any(s.startswith("Budget B") for s in seen), seen


def test_compare_names_the_budget_that_could_not_be_read(page, base, fake_ai, tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 nothing here")
    go(page, base, "compare")
    drop(page, "#compare-upload-a", GREEN)
    page.wait_for_selector("#compare-file-a .file-chip")
    drop(page, "#compare-upload-b", bad)
    page.wait_for_selector("#compare-file-b .file-chip")
    page.click("#compare-run")
    page.wait_for_selector(".notice-error", timeout=30_000)
    assert "Budget B" in page.inner_text(".notice-error")


# ── privacy and sessions ────────────────────────────────────────────────────
def test_privacy_is_stated_in_plain_words_on_both_tools(page, base):
    for name in ("evaluate", "compare"):
        go(page, base, name)
        assert ui.PRIVACY_TEXT in page.inner_text("main")


def sid_of(page):
    return page.get_attribute("#session-root", "data-sid")


def test_leaving_the_page_removes_the_visitors_data(page, base, fake_ai):
    run_evaluate(page, base)
    sid = sid_of(page)
    assert session.get(sid) is not None and session.current_job(sid, "evaluate") is not None
    page.evaluate("window.dispatchEvent(new Event('pagehide'))")
    assert wait_until(lambda: session.get(sid) is None)


def test_closing_the_tab_removes_the_visitors_data(browser, base, fake_ai):
    ctx = browser.new_context()
    pg = ctx.new_page()
    run_evaluate(pg, base)
    sid = sid_of(pg)
    pg.close(run_before_unload=True)
    assert wait_until(lambda: session.get(sid) is None, 8)
    ctx.close()


def test_refreshing_starts_clean(page, base, fake_ai):
    run_evaluate(page, base)
    old = sid_of(page)
    page.reload()
    page.wait_for_selector("#eval-upload")
    assert sid_of(page) != old and page.query_selector(".results") is None
    assert wait_until(lambda: session.get(old) is None)


def test_two_visitors_never_see_each_other(browser, base, fake_ai):
    a, b = browser.new_context().new_page(), browser.new_context().new_page()
    run_evaluate(a, base)
    go(b, base, "evaluate")
    time.sleep(1.5)
    assert b.query_selector(".results") is None and sid_of(a) != sid_of(b)
    a.context.close()
    b.context.close()


def test_the_session_routes_answer_quickly_and_ignore_nonsense(base):
    import urllib.request
    for route, body in (("ping", b"unknown"), ("end", b"unknown"), ("end", b"")):
        req = urllib.request.Request("%s/_session/%s" % (base, route), data=body, method="POST")
        with urllib.request.urlopen(req, timeout=5) as r:
            assert r.status == 204


def test_the_ping_keeps_a_session_alive_and_end_removes_it(base):
    import urllib.request
    sid = session.new_id()
    session.ensure(sid).last_seen -= 500
    before = session.get(sid).last_seen
    for route in ("ping", "end"):
        urllib.request.urlopen(urllib.request.Request("%s/_session/%s" % (base, route), data=sid.encode(), method="POST"), timeout=5)
        if route == "ping":
            assert session.get(sid).last_seen > before
    assert session.get(sid) is None


# ── wording and layout ──────────────────────────────────────────────────────
@pytest.mark.parametrize("name", ["home", "evaluate", "compare", "methodology"])
def test_no_technical_wording_on_any_page(page, base, name):
    go(page, base, name)
    page.wait_for_timeout(300)
    text = visible_text(page)
    hit = TECHNICAL_WORDS.search(text)
    assert not hit, "%s: %s" % (name, hit and hit.group(0))


def test_no_technical_wording_in_the_results(page, base, fake_ai):
    run_compare(page, base)
    assert not TECHNICAL_WORDS.search(visible_text(page))
    run_evaluate(page, base)
    assert not TECHNICAL_WORDS.search(visible_text(page))


@pytest.mark.parametrize("name", ["home", "evaluate", "compare", "methodology"])
def test_pages_fit_a_phone_screen(browser, base, name):
    ctx = browser.new_context(viewport={"width": 390, "height": 844})
    pg = ctx.new_page()
    go(pg, base, name)
    pg.wait_for_timeout(300)
    assert pg.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1")
    ctx.close()


def test_results_fit_a_phone_screen(browser, base, fake_ai):
    ctx = browser.new_context(viewport={"width": 390, "height": 844})
    pg = ctx.new_page()
    run_evaluate(pg, base)
    pg.wait_for_timeout(800)
    assert pg.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1")
    ctx.close()


def wait_until(condition, seconds=5.0):
    end = time.time() + seconds
    while time.time() < end:
        if condition():
            return True
        time.sleep(0.05)
    return condition()
