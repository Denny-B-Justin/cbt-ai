"""utils/export.py: the Excel files people download."""

import io

import openpyxl
import pytest

from tagging.pipeline import analyse_document, compare_budgets
from tests.conftest import noop_progress
from utils import compare, export
from utils import constants as C


@pytest.fixture
def analyses(fake_ai, green_pdf, traditional_pdf):
    return (analyse_document(green_pdf, "g.pdf", noop_progress), analyse_document(traditional_pdf, "t.pdf", noop_progress))


def sheet_rows(data, name):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return [[c for c in row] for row in wb[name].iter_rows(values_only=True)]


def test_evaluation_workbook_has_four_sheets(analyses):
    a, _ = analyses
    wb = openpyxl.load_workbook(io.BytesIO(export.evaluation_workbook(a)))
    assert wb.sheetnames == ["About", "Summary", "Rio Markers", "EU Taxonomy"]


def test_line_sheets_hold_every_line_with_plain_headings(analyses):
    a, _ = analyses
    data = export.evaluation_workbook(a)
    rio, eu = sheet_rows(data, "Rio Markers"), sheet_rows(data, "EU Taxonomy")
    assert len(rio) == len(eu) == a.stats["lines"] + 1
    assert rio[0][:4] == ["Line", "Page", "Ministry", "Budget line"] and rio[0][4].startswith("Amount") and "Rio marker" in rio[0] and "Why" in rio[0]
    assert "Climate tag" in eu[0] and "Main objective" in eu[0] and "Decided by" in eu[0]
    banned = ("rio_", "eu_", "tag_method", "decided_by", "climate_tag")
    assert not any(str(h).startswith(banned) for h in rio[0] + eu[0] if h)


def test_amounts_are_numbers_and_markers_are_readable(analyses):
    a, _ = analyses
    rio = sheet_rows(export.evaluation_workbook(a), "Rio Markers")
    header = rio[0]
    amounts = [r[4] for r in rio[1:]]
    assert all(isinstance(v, (int, float)) or v is None for v in amounts)
    markers = {r[header.index("Rio marker")] for r in rio[1:]}
    assert markers <= {"%d · %s" % (k, v) for k, v in C.RIO_LABELS.items()}


def test_summary_sheet_matches_the_screen_numbers(analyses):
    a, _ = analyses
    rows = sheet_rows(export.evaluation_workbook(a), "Summary")
    values = {(r[0], r[1]): r[2] for r in rows[1:]}
    from utils.formatting import fmt_pct
    assert values[(C.KIND_LABELS[C.KIND_EU], "Green share")] == fmt_pct(a.eu_summary["green_share"])
    assert values[("General", "Budget lines assessed")] == a.rio_summary["lines"]


def test_about_sheet_tells_people_to_review(analyses):
    a, _ = analyses
    about = dict((r[0], r[1]) for r in sheet_rows(export.evaluation_workbook(a), "About")[1:])
    assert about["Country"] == "Testland" and "reviewed" in about["Note"] and about["Pages read"] == a.pages


def test_text_from_a_document_never_becomes_a_formula(analyses):
    a, _ = analyses
    a.rio.loc[0, "programme"] = "=HYPERLINK(\"http://example.com\",\"click\")"
    a.rio.loc[1, "rationale"] = "=1+1"
    data = export.evaluation_workbook(a)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    cells = [c for row in wb["Rio Markers"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(cells) == 2 and all(c.data_type == "s" for c in cells)


def test_comparison_workbook_ends_with_the_overall_score(analyses):
    a, b = analyses
    comparison = compare_budgets(a, b)
    rows = compare.text_rows(compare.build_rows(a, b), comparison)
    data = export.comparison_workbook(a, b, comparison, rows)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Comparison", "Scores"]
    table = sheet_rows(data, "Comparison")
    assert table[0][1].startswith("Budget A: ") and table[0][2].startswith("Budget B: ")
    assert table[-1] == ["Overall climate score (out of 100)", str(comparison.scores["A"]), str(comparison.scores["B"])]
    scores = sheet_rows(data, "Scores")
    assert scores[1][1] == comparison.scores["A"] and scores[2][1] == comparison.scores["B"]


@pytest.mark.parametrize("name,stem", [("Georgia Budget 2024.pdf", "Georgia_Budget_2024pdf"), ("../../etc/passwd", "etcpasswd"),
                                       ("", "budget"), ("???", "budget"), ("a" * 100, "a" * 60)])
def test_file_names_are_safe(name, stem):
    assert export.file_stem(name) == stem
