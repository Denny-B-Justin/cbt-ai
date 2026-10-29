"""Builds the fictional sample budgets in this folder. Needs reportlab and python-docx (see requirements-dev.txt)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).resolve().parent
NOTE = "FICTIONAL SAMPLE - invented numbers for demonstration only. Not a real budget."

GREEN_FOCUS = {
    "file": "sample_budget_green_focus",
    "title": "Republic of Examplia - Ministry of Energy and Environment",
    "subtitle": "Programme Budget 2025 (amounts in EXM million)",
    "intro": ["The ministry plans to expand clean energy, protect communities from floods and improve water services in 2025."],
    "lines": [
        ("0101", "General administration", "Salaries, rent and office running costs of the ministry", 18.5),
        ("0201", "Solar power programme", "Construction of two 40 MW solar power plants and grid connection", 120.0),
        ("0202", "Wind energy development", "Wind farm feasibility studies and first turbines", 45.0),
        ("0203", "Household energy efficiency", "Insulation and efficient lighting for public housing", 30.0),
        ("0301", "Flood protection works", "River embankments and early warning systems against floods", 64.0),
        ("0302", "Drought-resistant farming", "Support for irrigation modernisation and drought-resistant seed varieties", 22.0),
        ("0401", "Wastewater treatment", "New municipal wastewater treatment plants", 58.0),
        ("0402", "Water supply networks", "Repair of leaking drinking water pipes", 36.0),
        ("0501", "National parks and wildlife", "Management of protected areas and habitat restoration", 17.0),
        ("0502", "Reforestation", "Planting trees on degraded land", 14.0),
        ("0601", "Waste recycling", "Recycling centres and composting of municipal waste", 21.0),
        ("0701", "Coal power station upgrade", "Life extension of the Northern coal-fired power station", 40.0),
        ("0801", "Road maintenance", "Resurfacing of regional roads", 35.0),
        ("0901", "Electric public transport", "Purchase of 60 electric buses for the capital", 52.0),
        ("1001", "Pensions and salaries", "Pensions and salaries of agency staff", 26.0),
        ("1101", "Statistics and audit", "Environmental statistics and audit services", 8.0),
    ],
}

TRADITIONAL = {
    "file": "sample_budget_traditional",
    "title": "Republic of Examplia - Ministry of Infrastructure and Social Services",
    "subtitle": "Programme Budget 2025 (amounts in EXM million)",
    "intro": ["The ministry focuses on roads, schools and hospitals in 2025."],
    "lines": [
        ("0101", "General administration", "Salaries, rent and office running costs of the ministry", 40.0),
        ("0201", "National roads", "Construction of a new four-lane highway", 210.0),
        ("0202", "Rural roads", "Rehabilitation of rural roads including drainage to withstand floods", 60.0),
        ("0301", "Primary schools", "Construction and equipment of primary schools", 120.0),
        ("0302", "Teachers' salaries", "Salaries of teachers", 150.0),
        ("0401", "District hospitals", "Construction of two district hospitals", 95.0),
        ("0402", "Medicines and supplies", "Purchase of essential medicines", 55.0),
        ("0501", "Social cash transfers", "Cash transfers to poor households", 85.0),
        ("0601", "Oil and gas exploration", "Support for oil and gas exploration licences", 30.0),
        ("0701", "Rooftop solar for schools", "Solar panels on 40 public school buildings", 12.0),
        ("0801", "Ports and harbours", "Expansion of the container port", 70.0),
        ("0901", "Police and justice", "Police stations and court buildings", 45.0),
        ("1001", "Pensions", "Pensions for retired civil servants", 66.0),
    ],
}

PROJECT = {
    "file": "sample_project_budget",
    "title": "Coastal Resilience Project (Examplia)",
    "subtitle": "Project budget summary (amounts in USD million)",
    "intro": ["A small project that protects coastal towns and restores wetlands."],
    "lines": [
        ("C1", "Sea walls and drainage", "Construction of sea walls and storm drainage to protect coastal towns from sea-level rise", 18.0),
        ("C2", "Mangrove and wetland restoration", "Restoring mangroves and wetlands that buffer storms", 7.5),
        ("C3", "Early warning system", "Cyclone and flood early warning stations", 3.2),
        ("C4", "Community training", "Training of local committees on disaster preparedness", 1.8),
        ("C5", "Project management", "Salaries and office costs of the project unit", 2.5),
    ],
}


def _total(lines):
    return sum(x[3] for x in lines)


def make_pdf(spec: dict) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(HERE / (spec["file"] + ".pdf")), pagesize=A4, title=spec["title"], leftMargin=36, rightMargin=36)
    body = [Paragraph(spec["title"], styles["Title"]), Paragraph(spec["subtitle"], styles["Heading3"]),
            Paragraph("<i>%s</i>" % NOTE, styles["Normal"]), Spacer(1, 10)]
    body += [Paragraph(t, styles["Normal"]) for t in spec["intro"]] + [Spacer(1, 12)]
    cell = styles["BodyText"]
    rows = [["Code", "Programme", "Description", "Amount"]]
    rows += [[c, Paragraph(p, cell), Paragraph(d, cell), "{:,.1f}".format(a)] for c, p, d, a in spec["lines"]]
    rows.append(["", Paragraph("<b>Total</b>", cell), "", "{:,.1f}".format(_total(spec["lines"]))])
    table = Table(rows, colWidths=[44, 130, 280, 60], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F3B2A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ALIGN", (3, 0), (3, -1), "RIGHT")]))
    body.append(table)
    doc.build(body)


def make_docx(spec: dict) -> None:
    doc = Document()
    doc.add_heading(spec["title"], 0)
    doc.add_heading(spec["subtitle"], 2)
    doc.add_paragraph(NOTE).italic = True
    for t in spec["intro"]:
        doc.add_paragraph(t)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Code", "Component", "Description", "Amount"]):
        table.rows[0].cells[i].text = h
    for c, p, d, a in spec["lines"]:
        cells = table.add_row().cells
        for i, v in enumerate([c, p, d, "{:,.1f}".format(a)]):
            cells[i].text = v
    cells = table.add_row().cells
    cells[1].text, cells[3].text = "Total", "{:,.1f}".format(_total(spec["lines"]))
    doc.save(str(HERE / (spec["file"] + ".docx")))


if __name__ == "__main__":
    make_pdf(GREEN_FOCUS)
    make_pdf(TRADITIONAL)
    make_pdf(PROJECT)
    make_docx(PROJECT)
    print("Samples written to", HERE)
